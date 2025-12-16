import pandas as pd
import os
import pdb
from openpyxl import load_workbook
from openpyxl.chart import BarChart3D, Reference, AreaChart, AreaChart3D, Series, LineChart, LineChart3D
import datetime
import win32com.client as win32
from win32com.client import Dispatch
from PIL import ImageGrab
import win32com.client       # Need pywin32 from pip
import pandas
import openpyxl
import excel2img
from openpyxl.chart.shapes import GraphicalProperties
import xlsxwriter


from openpyxl import Workbook
from openpyxl import drawing
from openpyxl.chart import (
    PieChart,
    ProjectedPieChart,
    Reference
)
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice



def pivot_attr(dataframe, kpi):
    table = pandas.DataFrame()

    table['Date'] = dataframe['Date']
    table['Total'] = dataframe[kpi]
    # table[kpi] = dataframe[kpi]
    writer = pd.ExcelWriter(str(os.getcwd()) + '\\Graphs\\output.xlsx', engine='xlsxwriter')
    table.to_excel(writer, 'index_', index=False)
    workbook_xlwt = writer.book
    worksheet_xlwt = writer.sheets['index_']


    #print('Here is the Graphs formation function!')

    chart_ = workbook_xlwt.add_chart({'type': 'line'})
    chart_.set_style(42)

    row_count_max = len(table)
    col_count_max = len(list(table.columns.values))


    chart_.add_series({
        'name': ['index_', 0, 1],
        'categories': ['index_', 1, 0, row_count_max, 0],
        'values': ['index_', 1, 1, row_count_max, 1],
        'marker': {'type': 'circle', 'size': 6},
        # 'line':       {'none': True},
        'line': {'width': 2.75}

    })
    chart_.set_chartarea({
        'border': {'none': True},
        # 'fill':   {'color': '#434343'},
        'gradient': {'colors': ['#000000', '#434343'], 'type': 'radial'}
        # 'font': {'color': 'white'}
        # 'name_font':{'name':'Calibri(Body)','size':14, 'color': 'white'}
    })
    chart_.set_plotarea({
        # 'border': {'color': 'red', 'width': 2, 'dash_type': 'dash'},
        'border': {'none': True},
        # 'fill':   {'color': '#434343'},
        'gradient': {'colors': ['#000000', '#434343'], 'type': 'radial'}
        # 'font': {'color': 'white'}
        # 'name_font':{'name':'Calibri(Body)','size':14, 'color': 'white'}
    })

    chart_.set_title({'name': kpi})
    #worksheet_xlwt.insert_chart("I3", chart_, {'x_offset': 25, 'y_offset': 10})
    worksheet_xlwt.insert_chart("I3", chart_, {'x_scale': 1.5, 'y_scale': 1.5})
    writer.save()
    workbook_xlwt.close()

    excel = win32com.client.Dispatch("Excel.Application")
    workbook = excel.Workbooks.Open(str(os.getcwd()) + "\\Graphs\\output.xlsx")

    wb_folder = workbook.Path
    wb_name = workbook.Name
    wb_path = os.path.join(wb_folder, wb_name)
    print("Processing Site Level KPI %s" % kpi)

    image_no = 0
    # import pdb
    # pdb.set_trace()

    for sheet in workbook.Worksheets:
        for n, shape in enumerate(sheet.Shapes):
            if shape.Name.startswith("Chart"):
                image_no += 1

                filename = kpi + ".png"
                file_path = os.path.join (wb_folder, filename)

                shape.Copy()  # Copies from Excel to Windows clipboard

                image = ImageGrab.grabclipboard()
                image.save(file_path, 'png')

    import time
    time.sleep(2)
    excel.Quit()